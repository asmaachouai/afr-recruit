"""
Text extraction from CV files.
Supports PDF (pdfplumber), DOCX (python-docx), and plain text.
"""

import io
from pathlib import Path

import pdfplumber
import structlog
from docx import Document

logger = structlog.get_logger(__name__)


class TextExtractor:
    """Extracts raw text from uploaded CV files."""

    SUPPORTED_TYPES = {
        "application/pdf": "_extract_pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "_extract_docx",
        "application/msword": "_extract_docx",
        "text/plain": "_extract_txt",
    }

    def extract(self, file_bytes: bytes, mime_type: str, filename: str) -> str:
        """
        Extract text from file bytes.
        Returns clean text string.
        """
        method_name = self.SUPPORTED_TYPES.get(mime_type)

        if not method_name:
            # Try to infer from extension
            ext = Path(filename).suffix.lower()
            if ext == ".pdf":
                method_name = "_extract_pdf"
            elif ext in (".docx", ".doc"):
                method_name = "_extract_docx"
            elif ext == ".txt":
                method_name = "_extract_txt"
            else:
                raise ValueError(f"Unsupported file type: {mime_type} / {ext}")

        method = getattr(self, method_name)
        text = method(file_bytes)

        logger.info(
            "Text extracted",
            filename=filename,
            mime_type=mime_type,
            char_count=len(text),
        )
        return text

    def _extract_pdf(self, file_bytes: bytes) -> str:
        """Extract text from PDF using pdfplumber."""
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
        return "\n".join(text_parts)

    def _extract_docx(self, file_bytes: bytes) -> str:
        """Extract text from DOCX preserving paragraph structure."""
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text.strip())
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)

    def _extract_txt(self, file_bytes: bytes) -> str:
        """Extract text from plain text file with encoding detection."""
        import chardet
        detected = chardet.detect(file_bytes)
        encoding = detected.get("encoding") or "utf-8"
        try:
            return file_bytes.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            return file_bytes.decode("utf-8", errors="replace")